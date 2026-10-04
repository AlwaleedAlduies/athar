import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_report():
    doc = Document()

    # Define Colors
    CLR_PRIMARY_HEX = "123B3B"       # Deep Emerald / Teal
    CLR_PRIMARY = RGBColor(18, 59, 59)
    CLR_GOLD_HEX = "C59B27"          # Heritage Gold
    CLR_GOLD = RGBColor(197, 155, 39)
    CLR_DARK_HEX = "0D2727"
    CLR_DARK = RGBColor(13, 39, 39)
    CLR_SAND_HEX = "F6F3EC"          # Parchment / Warm Sand
    CLR_BORDER_HEX = "D9D2C5"
    CLR_TEXT_DARK = RGBColor(30, 41, 59)
    CLR_TEXT_MUTED = RGBColor(100, 116, 139)
    CLR_WHITE = RGBColor(255, 255, 255)

    FONT_ARABIC = "Traditional Arabic"
    FONT_HEADING = "Traditional Arabic"

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69) # A4

    # Helper: set RTL on paragraph
    def set_rtl(p):
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def set_ltr(p):
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '0')
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def add_para(text="", style=None, space_before=0, space_after=6, line_spacing=1.2, align=WD_ALIGN_PARAGRAPH.RIGHT):
        p = doc.add_paragraph(style=style)
        set_rtl(p)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = align
        if text:
            run = p.add_run(text)
            format_run(run)
        return p

    def format_run(run, font_name=FONT_ARABIC, size_pt=13, color=CLR_TEXT_DARK, bold=False, italic=False):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rPr.append(rFonts)

    def add_heading_1(text):
        p = add_para(space_before=16, space_after=8)
        # Decorative colored mark
        r_mark = p.add_run("❖  ")
        format_run(r_mark, FONT_HEADING, 18, CLR_GOLD, bold=True)
        r_text = p.add_run(text)
        format_run(r_text, FONT_HEADING, 19, CLR_PRIMARY, bold=True)
        
        # Add bottom border via XML to paragraph
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{CLR_GOLD_HEX}"/></w:pBdr>')
        pPr.append(pBdr)
        return p

    def add_heading_2(text):
        p = add_para(space_before=12, space_after=6)
        r_mark = p.add_run("▪ ")
        format_run(r_mark, FONT_HEADING, 15, CLR_GOLD, bold=True)
        r_text = p.add_run(text)
        format_run(r_text, FONT_HEADING, 15, CLR_PRIMARY, bold=True)
        return p

    def add_heading_3(text):
        p = add_para(space_before=8, space_after=4)
        r_text = p.add_run(text)
        format_run(r_text, FONT_HEADING, 13.5, CLR_GOLD, bold=True)
        return p

    def add_bullet(bold_prefix, text):
        p = add_para(space_before=2, space_after=3)
        r_bullet = p.add_run("  • ")
        format_run(r_bullet, FONT_ARABIC, 12, CLR_GOLD, bold=True)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + ": ")
            format_run(r_bold, FONT_ARABIC, 12.5, CLR_PRIMARY, bold=True)
        r_desc = p.add_run(text)
        format_run(r_desc, FONT_ARABIC, 12, CLR_TEXT_DARK)
        return p

    def add_callout(title, body_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.6)
        
        # XML styling: background shading + thick gold right border
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CLR_SAND_HEX}"/>')
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="single" w:sz="36" w:space="0" w:color="{CLR_GOLD_HEX}"/>
            </w:tcBorders>
        ''')
        tcPr.append(shd)
        tcPr.append(borders)
        
        # Margins inside cell
        mar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="120" w:type="dxa"/>
                <w:bottom w:w="120" w:type="dxa"/>
                <w:left w:w="180" w:type="dxa"/>
                <w:right w:w="180" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(mar)

        cp = cell.paragraphs[0]
        set_rtl(cp)
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(2)
        r1 = cp.add_run(title)
        format_run(r1, FONT_HEADING, 13, CLR_PRIMARY, bold=True)

        bp = cell.add_paragraph()
        set_rtl(bp)
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(4)
        r2 = bp.add_run(body_text)
        format_run(r2, FONT_ARABIC, 12, CLR_TEXT_DARK, italic=True)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)

    def add_styled_image(img_path, caption):
        if not os.path.exists(img_path):
            print(f"Warning: image {img_path} not found.")
            return
        
        # Container table for bordered frame
        tbl = doc.add_table(rows=2, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Cell 0: Image
        c0 = tbl.cell(0, 0)
        c0.width = Inches(6.5)
        tcPr0 = c0._tc.get_or_add_tcPr()
        borders0 = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
                <w:left w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
            </w:tcBorders>
        ''')
        tcPr0.append(borders0)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(4)
        p0.paragraph_format.space_after = Pt(2)
        run0 = p0.add_run()
        run0.add_picture(img_path, width=Inches(6.2))
        
        # Cell 1: Caption bar
        c1 = tbl.cell(1, 0)
        c1.width = Inches(6.5)
        tcPr1 = c1._tc.get_or_add_tcPr()
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CLR_SAND_HEX}"/>')
        borders1 = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:color="{CLR_GOLD_HEX}"/>
                <w:left w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
                <w:bottom w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
                <w:right w:val="single" w:sz="8" w:color="{CLR_BORDER_HEX}"/>
            </w:tcBorders>
        ''')
        tcPr1.append(shd1)
        tcPr1.append(borders1)
        
        p1 = c1.paragraphs[0]
        set_rtl(p1)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(4)
        p1.paragraph_format.space_after = Pt(4)
        r_cap_label = p1.add_run("شكل توضيحي: ")
        format_run(r_cap_label, FONT_HEADING, 11, CLR_GOLD, bold=True)
        r_cap = p1.add_run(caption)
        format_run(r_cap, FONT_ARABIC, 11, CLR_TEXT_DARK, italic=True)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(8)

    # ----------------------------------------------------
    # COVER PAGE
    # ----------------------------------------------------
    # Top spacing
    p_top = add_para(space_before=18, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Platform Category
    p_pre = add_para("وثيقة الإنجاز والتقرير الوظيفي الشامل للمشروع", space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_run(p_pre.runs[0], FONT_ARABIC, 13, CLR_GOLD, bold=True)

    # Main Title
    p_title = add_para("مـنـصـة أَثَــر | ATHAR", space_before=4, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_run(p_title.runs[0], FONT_HEADING, 32, CLR_PRIMARY, bold=True)

    # Subtitle
    p_sub = add_para("استكشاف التاريخ الإسلامي والسيرة النبوية بالمعرفة المسندة", space_before=2, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_run(p_sub.runs[0], FONT_ARABIC, 16, CLR_GOLD, bold=True)

    # Slogan Banner Box
    add_callout("شعار المنصة وفلسفتها التأسيسية:", "« لا تقرأ التاريخ فقط ... عِش أثره »\nالذكاء الاصطناعي ليس المصدر — التاريخ يُستقى من نصوصه وأدلته الموثقة.")

    # Heritage banner image
    if os.path.exists("static/img/heritage.png"):
        p_img = add_para(space_before=6, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_img = p_img.add_run()
        r_img.add_picture("static/img/heritage.png", width=Inches(6.4))

    # Meta table at bottom of cover
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("إصدار الوثيقة والمنصة:", "الإصدار 1.0 (جاهز للتشغيل والإنتاج)"),
        ("تاريخ إعداد التقرير:", "أكتوبر 2026 م"),
        ("نطاق العمليات:", "دليل شامل لكافة العمليات الوظيفية الجاهزة للمستكشف والإدارة"),
        ("مستودع الكود المصدري:", "https://github.com/AlwaleedAlduies/athar"),
    ]
    for row_idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        for c in (c0, c1):
            tcPr = c._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CLR_SAND_HEX}"/>')
            bdr = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:bottom w:val="single" w:sz="4" w:color="{CLR_BORDER_HEX}"/>
                    <w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(shd)
            tcPr.append(bdr)
        
        p0 = c0.paragraphs[0]
        set_rtl(p0)
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        format_run(r0, FONT_HEADING, 11, CLR_PRIMARY, bold=True)
        
        p1 = c1.paragraphs[0]
        set_rtl(p1)
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(val)
        format_run(r1, FONT_ARABIC, 11, CLR_TEXT_DARK)

    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 1: VISION & DEFINITION
    # ----------------------------------------------------
    add_heading_1("1. التعريف بالمنصة، الرؤية، والرسالة الإستراتيجية")
    
    add_heading_2("ما هي منصة أثَر؟")
    add_para("منصة «أثَر» هي بيئة رقمية تفاعلية متقدمة ومتخصصة في استكشاف التاريخ الإسلامي وأحداث السيرة النبوية الشريفة. تنطلق المنصة من مبدأ أصيل يجمع بين جاذبية السرد القصصي المعاصر والصرامة الأكاديمية والتوثيق المصدري الصارم، متجاوزةً بذلك عيوب القراءة الخطية الورقية ومحاذير الاعتماد غير المنضبط على الذكاء الاصطناعي التوليدي.")

    add_heading_2("الرؤية (Vision)")
    add_para("أن تكون منصة «أثَر» المرجع الرقمي التفاعلي الأول عالميًا في عرض التاريخ والتراث الإسلامي، من خلال تقديم تجربة بصرية ومعرفية رائدة تجعل المتلقي يعايش وقائع التاريخ ويفهم سياقاته وقراراته ومآلاته، مع إتاحة التحقق الفوري المباشر من كل معلومة ودور من مصادرها الأصلية المعترف بها.")

    add_heading_2("الرسالة (Mission)")
    add_para("إعادة بعث التراث التاريخي الإسلامي بأسلوب عصري جذاب عبر شبكة معرفية متكاملة (Knowledge Graph) تربط بين الشخصيات والأحداث والأماكن والمصادر، وترسيخ قاعدة: «الذكاء الاصطناعي ليس المصدر؛ بل هو خادم ومحلل للنصوص، والأدلة تستند حصرًا إلى مقتطفاتها الأصلية».")

    add_heading_2("الأهداف الإستراتيجية للمشروع")
    add_bullet("التوثيق المصدري اللحظي", "ربط كل فقرة سردية أو ادعاء بمقتطف حرفي ورقم صفحة واسم مرجع ورابط مباشر على الإنترنت.")
    add_bullet("الفصل المنهجي الصارم للمعلومات", "تصنيف كل معلومة تاريخية بدقة إلى: (حقيقة موثقة، تفسير تحليلي، اختلاف روايات، أو مسألة غير محسومة)، وتجنب القطع فيما لم يثبت.")
    add_bullet("التعلم التاريخي السببي", "استيعاب السيرة النبوية بوصفها حركة حية لها مقدمات وأسباب وقرارات بشرية ونتائج ومآلات مدروسة.")
    add_bullet("عزل المحاكاة التوعوية", "توفير أنشطة استكشافية افتراضية من طراز «ماذا لو؟» بصورة معزولة تمامًا ومصرح بها لمنع اختلاط الافتراض بالرواية الثابتة.")
    add_bullet("أمانة المعرفة والتمكين المؤسسي", "تزويد المحققين والمحررين («أمناء المعرفة») ببوابة إدارة مستقلة متكاملة لتحليل المصادر واستخلاص الكيانات واعتماد النشر بأمان.")

    # ----------------------------------------------------
    # SECTION 2: ARCHITECTURE & METRICS
    # ----------------------------------------------------
    add_heading_1("2. المعمارية التقنية وشبكة المعرفة والإنجاز الرقمي")
    
    add_para("شُيدت المنصة كمنظومة برمجية حقيقية متكاملة للإنتاج الفعلي (Production-Ready) وفق أحدث المعايير:")
    add_bullet("العمود الفقري", "مبني على Django 5.2 وDjango REST Framework (DRF) على بيئة Python 3.12+.")
    add_bullet("قاعدة البيانات", "مهيأة للعمل الفوري بـ SQLite3 محليًا، ومجهزة للإنتاج المباشر مع PostgreSQL عبر Docker وWaitress.")
    add_bullet("الواجهة الأمامية", "واجهة عربية أصيلة بنمط اتجاه اليمين إلى اليسار (RTL)، مع أداء سريع جدًا واعتماد على Vanilla CSS وModern JS دون أطر واجهات ثقيلة.")
    add_bullet("محرك الاسترجاع المتجهي (RAG)", "محرك هجين يعتمد التضمين المعجمي المحلي (Lexical Hash) أو التضمين الدلالي عبر Ollama، ومحكوم بصرامة حصر الأدلة داخل نطاق الحدث لمنع الهلوسة والتداخل.")

    add_heading_2("إحصائيات الإنجاز الفعلي المثبت في قاعدة البيانات")
    add_para("لا تعتمد المنصة على محتوى وهمي أو تجريبي، بل تتضمن مكتبة محققة وموثقة بالكامل وفق الجدول التالي:")

    # Table of DB metrics
    metrics_table = doc.add_table(rows=9, cols=3)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["المكون / الكيان", "العدد المنجز والمثبت", "الوصف الوظيفي والتأكيد"]
    data_metrics = [
        ("الأحداث التاريخية المنشورة", "6 محطات كبرى", "الهجرة، بدر، أُحد، الخندق، الحديبية، فتح مكة"),
        ("فقرات السرد التفاعلي", "46 فقرة معتمدة", "موزعة عبر 29 فصلاً بمجموع 2,692 كلمة سرد متصل"),
        ("الادعاءات والمعلومات المصنفة", "107 ادعاءات", "مصنفة بين حقائق وتفسيرات وروايات متعددة"),
        ("الأدلة النصية الحرفية", "133 دليلاً", "مطابقة لنصوص المصادر بنسبة 100% مع أرقام الصفحات والمواضع"),
        ("علاقات الشبكة المعرفية", "106 علاقات", "مشاركة، رواية خبر، أسباب، نتائج، توقيت، وأماكن"),
        ("الكيانات الكلية المسجلة", "72 كياناً مسجلاً", "32 شخصية موثقة الأدوار، 6 أماكن تاريخية، 27 مصدراً، حقبة واحدة"),
        ("المقاطع المفهرسة بالمتجهات", "58 مقطعاً نصياً", "مقاطع مفهرسة ومجزأة مع تخزين المتجهات الرقمية للاسترجاع"),
        ("جودة الاختبارات البرمجية", "163 اختباراً ناجحاً", "اجتياز اختبارات Django وDRF والأمان بنسبة 100% وبلا أخطاء"),
    ]

    # Style Table Header
    hdr_row = metrics_table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CLR_PRIMARY_HEX}"/>')
        tcPr.append(shd)
        p = cell.paragraphs[0]
        set_rtl(p)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(heading)
        format_run(run, FONT_HEADING, 11.5, CLR_WHITE, bold=True)

    # Style Table Body
    for row_idx, (col0, col1, col2) in enumerate(data_metrics, 1):
        row = metrics_table.rows[row_idx]
        bg_color = CLR_SAND_HEX if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate((col0, col1, col2)):
            cell = row.cells[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            bdr = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:bottom w:val="single" w:sz="4" w:color="{CLR_BORDER_HEX}"/>
                    <w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(shd)
            tcPr.append(bdr)
            p = cell.paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(val)
            is_bold = (c_idx == 0 or c_idx == 1)
            clr = CLR_PRIMARY if c_idx == 1 else CLR_TEXT_DARK
            format_run(run, FONT_ARABIC, 11, clr, bold=is_bold)

    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 3: FUNCTIONAL OPERATIONS REPORT
    # ----------------------------------------------------
    add_heading_1("3. التقرير الوظيفي الشامل للعمليات الجاهزة في النظام")
    add_para("يستعرض هذا القسم بالتفصيل كافة العمليات والوظائف التي يستطيع المستخدمون وأمناء المعرفة القيام بها فعليًا داخل المنصة:")

    add_heading_2("أولاً: مسار وعمليات الزائر والمستكشف (Explorer Operations)")
    
    add_heading_3("1. إدارة الحساب والتخصيص الشخصي")
    add_bullet("إنشاء حساب مستكشف جديد", "واجهة تسجيل بسيطة تفرض شروط الأمان وكلمات المرور المعتمدة.")
    add_bullet("تسجيل الدخول والخروج", "جلسات مؤمنة بـ CSRF مع نظام حماية يمنع محاولات التخمين الآلية (30 محاولة/ساعة).")
    add_bullet("الملف الشخصي", "تعديل النبذة التعريفية والاسم والبريد.")
    add_bullet("المكتبة والمحفوظات الخاصة (/library/)", "إمكانية حفظ أي حدث أو شخصية أو مكان بالمفضلة بضغطة زر واحدة، مع سجل تلقائي للمحطات التي تم استكشافها، وسجل بعمليات البحث السابقة.")

    add_heading_3("2. تجربة «عِش الرحلة» التفاعلية (/journey/)")
    add_bullet("التنقل المتسلسل الفصلي", "قراءة المحطات وفق مسار سردي ثلاثي لكل محطة: (المشهد التمهيدي ➔ السياق والقرارات والنتائج ➔ أدلة الإسناد).")
    add_bullet("مؤشر الإنجاز البصري", "شريط تقدم ديناميكي أعلى الشاشة يحسب نسبة استكمال فصول المحطة وصولاً إلى 100%.")
    add_bullet("المسار المركز (مسار المدينة)", "تصفية الرحلة لتركز حصرًا على المحطات الكبرى التأسيسية في المدينة عبر المعامل ?path=medina.")
    add_bullet("مشاركة الفصول بروابط مباشرة", "إمكانية نسخ رابط مباشر ينقل القارئ إلى حدث وفصل محدد تلقائيًا مثل (#event=10&chapter=evidence).")
    add_bullet("سهولة الوصول والتنقل", "دعم كامل للتنقل بالأسهم في لوحة المفاتيح وفق الاتجاه العربي RTL، وزر خاص لحفظ تفضيل تقليل الحركة البصرية (Reduced Motion).")

    add_heading_3("3. السرد التفاعلي والاستشهاد الموضعي بنقرة واحدة (In-Text Tracing)")
    add_bullet("فتح درج الاستشهاد الحرفي", "النقر على أي موضع من الفقرة السردية يفتح نافذة جانبية فورية تعرض:")
    add_bullet("تفاصيل المرجع", "نص المقتطف الحرفي المستند إليه، اسم الكتاب أو المصدر، رقم الصفحة الفعلي، ورابط المرجع الأصلي على الإنترنت.")
    add_bullet("اختصارات لوحة المفاتيح", "استخدام مفتاحي Enter والمسافة لفتح الاستشهاد، ومفتاح Escape لإغلاقه مع إعادة مؤشر التركيز إلى الفقرة.")

    add_heading_3("4. الخط الزمني التاريخي المرشح (/timeline/)")
    add_bullet("العرض الزمني المزدوج", "ترتيب المحطات وفق التسلسل التاريخي بالسنوات الهجرية والميلادية.")
    add_bullet("التصفية متعددة المستويات", "ترشيح الأحداث بحسب العصور الإسلامية، أو حسب نوع الحدث (معركة، صلح وعهد، رحلة وتحول، علم وحضارة).")
    add_bullet("البحث السريع المدمج", "تصفية الخط الزمني نصيًا مع دعم الترقيم بـ 18 حدثاً لكل صفحة.")

    add_heading_3("5. صفحات الكيانات المعرفية المعمقة (/explore/<id>/)")
    add_bullet("صفحة الحدث التاريخي", "عرض بطاقات تفصيلية للسياق، والأسباب، والقرارات، والنتائج، والمآلات.")
    add_bullet("خريطة العلاقات الشبكية (SVG Graph)", "رسم بياني تفاعلي لعلاقات الحدث بالأشخاص والأماكن، مع أزرار لإظهار أو إخفاء أنواع الكيانات بنقرة واحدة.")
    add_bullet("بطاقات أدوار الشخصيات", "عرض الشخصيات المشاركة مع توضيح الدور الدقيق المسند بدليله المعتمد (مع تمييز المشاركة عن مجرد رواية الخبر).")
    add_bullet("أدلة الأماكن والتنقل الزمني", "استعراض أدلة الموقع الجغرافي، وروابط سريعة للمحطة التاريخية السابقة والتالية.")
    add_bullet("صفحة الشخصية", "سيرة الشخصية، مع سجل زمني شامل يجمع كافة أدوارها في جميع الأحداث مدعمة بالأدلة والمصادر.")
    add_bullet("صفحة المكان", "الإحداثيات الجغرافية، الأهمية التاريخية، وسجل الوقائع التي شهدها المكان مدعومة بأدلتها.")
    add_bullet("صفحة المصدر", "بيانات الطبعة والمؤلف، قائمة المقاطع المفهرسة، وتحميل ملف المصدر الأصلي بأمان.")

    add_heading_3("6. محرك البحث العربي الموحد (/search/)")
    add_bullet("تطبيع النصوص العربية", "خوارزمية ذكية تتسامح مع كتابة الهمزات (أ، إ، آ)، والياء والألف اللينة (ي، ى)، والتاء المربوطة (ة، ه)، وإسقاط التشكيل.")
    add_bullet("التصنيف النوعي للنتائج", "تصفية البحث ليعرض حصرًا (أحداثاً، شخصيات، أماكن، مصادر، أو عصوراً).")

    add_heading_3("7. المرشد التاريخي الذكي («توقف واسأل» - AI RAG)")
    add_bullet("طرح الأسئلة الحرة داخل الحدث", "استفسار الزائر عن وقائع الحدث وأدواره ونتائجه أثناء القراءة.")
    add_bullet("إجابات مستندة للأدلة حصراً", "توليد شروح مسندة بالمقتطفات والمراجع، مع إبراز روايات الخلاف بدقة، والامتناع التام عن الجواب عند نقص الأدلة.")
    add_bullet("سجل الإجابات الشخصي (/library/answers/)", "أرشفة كافة أسئلة المستكشف السابقة وإجاباتها بمراجعها لسهولة المراجعة لاحقاً.")

    add_heading_3("8. أنشطة المحاكاة الافتراضية («ماذا لو؟»)")
    add_bullet("الموافقة والتأكيد الصريح", "اشتراط إقرار الزائر بفهم طبيعة المحاكاة الافتراضية قبل الدخول إليها.")
    add_bullet("تحليل نقاط القرار والبدائل", "استعراض القرارات التاريخية، والبديل الافتراضي، والنتائج المحتملة، والعوامل المؤثرة المعزولة تماماً عن ثوابت السرد.")

    # ----------------------------------------------------
    # SECTION 3B: ADMIN OPERATIONS
    # ----------------------------------------------------
    add_heading_2("ثانياً: مسار وعمليات أمين المعرفة ولوحة الإدارة (Dashboard & Editorial)")

    add_heading_3("1. لوحة المؤشرات والإشراف المركزي (/dashboard/)")
    add_bullet("الإحصائيات اللحظية", "متابعة أعداد المحتوى المنشور، والادعاءات التي بانتظار المراجعة، والمصادر غير المفهرسة، والأسئلة التي تتطلب تدقيقاً.")
    add_bullet("المراقبة السريعة", "جدول لآخر استفسارات الذكاء الاصطناعي لفحص كفاءة الردود وأزمنة الاستجابة.")

    add_heading_3("2. إدارة المحتوى والتحرير في سياقه (Contextual CMS)")
    add_bullet("التحكم الكامل (CRUD)", "إضافة وتعديل وحذف وأرشفة واستعادة لكافة الأصناف المعرفية الـ 11 في النظام.")
    add_bullet("مساحة التحرير السياقية", "إدارة فصول وقصص السرد والأدلة وعلاقات الشخصيات والأماكن مباشرة من داخل بطاقة الحدث دون تشتت.")
    add_bullet("البحث والفرز المتقدم", "تصفية السجلات حسب حالة النشر (مسودة، مراجعة، معتمد، منشور، مؤرشف) وحسب الحدث أو المصدر.")
    add_bullet("الربط السريع بالبحث", "مربعات اختيار ذكية تدعم البحث النصي لربط الشخصيات والأماكن والأدلة بسلاسة.")

    add_heading_3("3. الاستيعاب الذكي للمصادر والتحليل الذاتي الآلي (/dashboard/intake/)")
    add_bullet("الاستيراد متعدد الوسائط", "رفع ملفات PDF نصية (مع أرقام الصفحات حتى 500 صفحة)، مستندات DOCX، ملفات TXT، أو روابط صفحات ويب عامة.")
    add_bullet("التنقية الآلية لصفحات الويب", "عزل متن المقال والحواشي تلقائيًا واستبعاد القوائم والسكربتات والإعلانات مع حماية تامة ضد هجمات SSRF.")
    add_bullet("التحليل الذاتي بالذكاء الاصطناعي", "استخراج الكيانات المذكورة في المصدر (أشخاص، أماكن، أحداث) مقرونة باقتباساتها الحرفية ومواضعها.")
    add_bullet("اتخاذ القرار الفوري", "ربط الاسم المكتشف بسجل قائم كـ «رابط ذكر»، أو إنشاء مسودة سجل جديد، أو التجاهل.")
    add_bullet("التحليل الجزئي والاستكمال", "إمكانية إيقاف واستئناف تحليل المصادر الطويلة على دفعات دون تكرار المقاطع المكتملة.")

    add_heading_3("4. استوديو صياغة المعرفة والسرد (/dashboard/studio/)")
    add_bullet("انتقاء المقاطع الموجه", "اختيار حدث معين وتحديد من 1 إلى 12 مقطعاً نصياً من مصدر معتمد لتحليلها.")
    add_bullet("توليد مقترحات السرد والأدلة", "صياغة مسودات لفقرات القصة والمعلومات وأدوار الشخصيات المستندة حصريًا للمقاطع المحددة.")
    add_bullet("المراجعة والاعتماد البشري", "رفض الخادم لأي اقتباس لا يطابق النص حرفيًا، وإبقاء كافة المقترحات كمسودات بانتظار الاعتماد والمراجعة البشرية.")

    add_heading_3("5. اعتماد المصادر وتأمين الاستشهادات")
    add_bullet("الاعتماد الشامل بضغطة زر", "فهرسة النص، بناء المتجهات، ترقية حالة المصدر لمعتمد، وتفعيله لمحرك RAG في خطوة واحدة.")
    add_bullet("حماية المقاطع المستشهد بها", "منع النظام حذف أو تعديل نص أي مقطع ترتبط به أدلة أو استشهادات منشورة حفاظاً على الثبات التاريخي.")
    add_bullet("إصدارات المصادر (Revisions)", "حفظ النسخ المنقحة كمصادر جديدة ترتبط بالإصدار السابق عبر حقل replaces دون إفساد الاستشهادات القديمة.")

    add_heading_3("6. إدارة وتفعيل مزودي الذكاء الاصطناعي (/dashboard/ai-providers/)")
    add_bullet("دعم المزودات الأربعة", "تهيئة اتصالات Ollama (محلي)، Gemini، OpenAI، أو الخدمات المتوافقة مع OpenAI.")
    add_bullet("الفحص والاختبار الحي", "زر فحص يرسل طلباً تجريبياً للتحقق من استجابة النموذج وتوليد JSON سليم قبل التفعيل.")
    add_bullet("التفعيل بضغطة زر والتشفير", "تفعيل المزوّد الناجح فوراً مع تشفير المفاتيح باستخدام Fernet وحجبها عن واجهات التحرير والمتصفح.")

    add_heading_3("7. الرقابة والتدقيق التاريخي لسجلات الذكاء الاصطناعي")
    add_bullet("سجل الأسئلة الشامل", "مراجعة كافة أسئلة الزوار، والإجابات المولدة، والمراجع المسترجعة، وزمن الاستجابة بالميلي ثانية.")
    add_bullet("تصفية المراجعة البشرية", "فلترة الأسئلة التي صُنفت بأنها تحتاج مراجعة أو التي تعذر تقديم جواب كافٍ عنها.")

    add_heading_3("8. الأرشفة الناعمة والاستعادة")
    add_bullet("معاينة أثر الحذف", "نافذة معاينة توضح السجلات المرتبطة قبل الأرشفة أو الحذف.")
    add_bullet("الاستعادة كمسودة", "إمكانية استعادة أي حدث أو شخصية أو مصدر مؤرشف وإعادته كمسودة قابلة لإعادة التحرير والنشر.")

    doc.add_page_break()

    # ----------------------------------------------------
    # SECTION 4: VISUAL SCREENSHOTS GALLERY
    # ----------------------------------------------------
    add_heading_1("4. معرض لقطات الشاشة الحية للمنصة (Visual Interface Gallery)")
    add_para("تجسد هذه اللقطات المأخوذة مباشرة من النظام المحلي الجاري تشغيله الهوية البصرية المعتمدة وأبرز الواجهات والوظائف:")

    screenshots_data = [
        ("screenshots/screen_landing.png", "الواجهة الرئيسية لمنصة أثَر: إبراز الفلسفة التأسيسية، الشعار البصري، وأبرز محطات السيرة النبوية المحققة."),
        ("screenshots/screen_journey.png", "مسار «عِش الرحلة» التفاعلي: استعراض الفصول الثلاثة لكل محطة مع شريط تقدم القراءة وإمكانية الفرز لمسار المدينة."),
        ("screenshots/screen_event_hijrah.png", "صفحة حدث الهجرة النبوية: السرد القصصي، شبكة العلاقات التفاعلية، درج الاستشهاد الحرفي، وبطاقات أدوار الشخصيات والأماكن."),
        ("screenshots/screen_timeline.png", "الخط الزمني التاريخي: استعراض الأحداث بالهجري والميلادي مع فلاتر العصور وأنواع الوقائع (معارك، معاهدات، رحلات)."),
        ("screenshots/screen_discover.png", "مركز الاكتشاف والمكتبة: تصفح العصور، استئناف القراءة، وبطاقة المسار المتخصص لأحداث المدينة المنورة."),
        ("screenshots/screen_login.png", "بوابة تسجيل الدخول وأمانة المعرفة: واجهة آمنة ومحمية بـ CSRF لصلاحيات المستكشفين وأمناء المعرفة."),
    ]

    for img_path, caption in screenshots_data:
        add_styled_image(img_path, caption)

    # ----------------------------------------------------
    # SECTION 5: CONCLUSION & DEPLOYMENT
    # ----------------------------------------------------
    add_heading_1("5. الخاتمة والجاهزية للنشر والتشغيل الفوري")
    add_para("يمثل مشروع «أثَر» نموذجًا مكتمل الأركان لمنصة معرفية عربية تحترم أصول التوثيق التاريخي وتوظف أحدث ممارسات البرمجة والذكاء الاصطناعي لخدمة التراث الإسلامي.")

    add_heading_2("الجاهزية للتشغيل الفوري (Turnkey Status)")
    add_bullet("المستودع السحابي", "المشروع بالكامل مرفوع على GitHub: https://github.com/AlwaleedAlduies/athar.")
    add_bullet("قاعدة البيانات المدرجة", "قاعدة البيانات db.sqlite3 مدمجة بالمستودع بكامل بياناتها المحققة بعد تأمينها وإزالة المفاتيح السرية.")
    add_bullet("حساب الإدارة الافتراضي", "مستخدم الإدارة جاهز وموثق: اسم المستخدم admin | كلمة المرور admin123456.")
    add_bullet("التشغيل بضغطة زر واحدة", "يمكن لأي مستخدم تشغيل المنصة فورًا عبر .\\start.ps1 على Windows أو ./start.sh على Linux/macOS.")

    add_heading_2("الخطوات التوسعية المستقبلية المقترحة")
    add_bullet("توسيع الحقب التاريخية", "تغذية عصر الخلفاء الراشدين والدولة الأموية والعباسية بنفس منهجية السرد المسند عبر استوديو المعرفة.")
    add_bullet("النشر السحابي عالي الإتاحة", "ربط قاعدة بيانات PostgreSQL على Render أو سيرفر VPS مع إعداد HTTPS وخدمات البريد.")
    add_bullet("طابور المهام غير المتزامن", "دمج Celery / Redis لمعالجة ملفات PDF الضخمة في الخلفية دون أي تأخير في استجابة الواجهة.")

    # Save document
    output_path = r"c:\Users\LOQ\Documents\Work\منصة اثر\athar\تقرير_مشروع_منصة_أثر_الشامل.docx"
    doc.save(output_path)
    print(f"Report successfully generated at: {output_path}")

if __name__ == "__main__":
    create_report()
