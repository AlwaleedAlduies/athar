import io
import json
import os
import tempfile
from unittest.mock import patch
from django.test import TestCase, Client, override_settings
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from rest_framework.test import APIClient
from apps.accounts.models import User, Bookmark, ExplorationHistory
from apps.history.models import Entity, HistoricalEra, HistoricalEvent, HistoricalPerson, HistoricalPlace, Status
from apps.history.search import normalize, search_entities
from apps.sources.models import HistoricalSource, SourceChunk
from apps.sources.services import process_source
from apps.sources.validators import validate_document
from apps.knowledge.models import HistoricalClaim, Evidence, EntityRelationship
from apps.knowledge.selectors import graph_for, trace_claim
from apps.ai.services import ask, retrieve, parse_answer, NO_EVIDENCE
from apps.ai.models import AIQueryLog
from apps.ai.embeddings import HashEmbedding
from apps.simulations.models import SimulationScenario
from apps.dashboard.forms import editor_form

@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class PlatformTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user(username='curator', password='Test!Passphrase928', role='ADMIN')
        cls.user = User.objects.create_user(username='reader', password='Test!Passphrase928')
        cls.other = User.objects.create_user(username='reader2', password='Test!Passphrase928')
        cls.era = HistoricalEra.objects.create(title='حقبة اختبار', slug='test-era', status='PUBLISHED')
        cls.place = HistoricalPlace.objects.create(title='أرض الاختبار', slug='test-place', status='PUBLISHED')
        cls.event = HistoricalEvent.objects.create(title='حدث الاختبار', slug='test-event', era=cls.era, location=cls.place, status='PUBLISHED')
        cls.draft = HistoricalEvent.objects.create(title='مسودة خاصة', slug='private-event', era=cls.era)
        cls.source = HistoricalSource.objects.create(title='دليل اختبار برمجي', slug='test-source', status='PUBLISHED', is_approved=True)
        cls.event.sources.add(cls.source)
        cls.text = 'هذا نص اختبار برمجي لا يمثل معلومة تاريخية. توثيق القرار يتطلب ربط المعلومة بمصدر معتمد.'
        cls.chunk = SourceChunk.objects.create(source=cls.source, ordinal=0, text=cls.text, page_number=7, section='اختبارات الثقة', embedding=HashEmbedding().embed(cls.text), embedding_reference=HashEmbedding.reference)
        cls.claim = HistoricalClaim.objects.create(event=cls.event, claim_text='توثيق القرار يتطلب ربط المعلومة بمصدر معتمد.', claim_type='FACT', status='APPROVED')
        cls.evidence = Evidence.objects.create(claim=cls.claim, source=cls.source, source_chunk=cls.chunk, evidence_text=cls.claim.claim_text, page_number=7, section='اختبارات الثقة')
        cls.scenario = SimulationScenario.objects.create(title='نشاط اختبار', event=cls.event, decision_point='قرار افتراضي', alternative='بديل محتمل', possible_consequences='احتمالات', influencing_factors='عوامل', status='PUBLISHED')

    def setUp(self):
        cache.clear()
        self.api = APIClient()
        self.environment = patch.dict(os.environ, {'LLM_PROVIDER':'extractive', 'EMBEDDING_BACKEND':'hash'})
        self.environment.start()
        self.addCleanup(self.environment.stop)

    def test_registration_cannot_escalate_role(self):
        response = self.client.post('/register/', {'username':'new-reader', 'first_name':'مستكشف', 'email':'reader@example.test', 'password1':'New!LongPhrase729', 'password2':'New!LongPhrase729', 'role':'ADMIN', 'is_superuser':True})
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username='new-reader')
        self.assertEqual(user.role, 'USER')
        self.assertFalse(user.is_superuser)
        self.assertTrue(user.check_password('New!LongPhrase729'))

    def test_login_logout_and_private_profile(self):
        self.assertEqual(self.client.get('/profile/').status_code, 302)
        self.assertTrue(self.client.login(username='reader', password='Test!Passphrase928'))
        self.assertEqual(self.client.get('/profile/').status_code, 200)
        self.assertEqual(self.client.get('/logout/').status_code, 405)
        self.assertEqual(self.client.post('/logout/').status_code, 302)

    def test_arabic_auth_even_with_english_browser(self):
        response = self.client.get('/login/', HTTP_ACCEPT_LANGUAGE='en-US,en;q=0.9')
        self.assertContains(response, 'اسم المستخدم')
        self.assertContains(response, 'كلمة المرور')
        self.assertNotContains(response, 'Username:')

    def test_regular_user_cannot_write_any_historical_resource(self):
        self.api.force_authenticate(self.user)
        for path in ['events', 'persons', 'places', 'eras', 'sources', 'claims', 'evidence', 'relationships', 'simulations']:
            with self.subTest(path=path):
                self.assertEqual(self.api.post(f'/api/v1/{path}/', {}, format='json').status_code, 403)
        self.assertEqual(self.api.patch(f'/api/v1/events/{self.event.pk}/', {'status':'APPROVED'}, format='json').status_code, 403)
        self.assertEqual(self.api.delete(f'/api/v1/events/{self.event.pk}/').status_code, 403)

    def test_dashboard_permission_enforced_on_server(self):
        self.client.force_login(self.user)
        for path in ['/dashboard/', '/dashboard/events/', '/dashboard/events/new/', '/dashboard/logs/', '/dashboard/users/']:
            self.assertEqual(self.client.get(path).status_code, 403)

    def test_visibility_detail_and_api(self):
        ids = [row['id'] for row in self.api.get('/api/v1/events/').data['results']]
        self.assertIn(self.event.pk, ids)
        self.assertNotIn(self.draft.pk, ids)
        self.assertEqual(self.api.get(f'/api/v1/events/{self.draft.pk}/').status_code, 404)
        self.assertEqual(self.client.get(self.draft.get_absolute_url()).status_code, 404)
        self.api.force_authenticate(self.admin)
        self.assertEqual(self.api.get(f'/api/v1/events/{self.draft.pk}/').status_code, 200)

    def test_event_crud_and_archive(self):
        self.api.force_authenticate(self.admin)
        response = self.api.post('/api/v1/events/', {'title':'حدث جديد', 'slug':'new-event', 'era':self.era.pk, 'status':'DRAFT'}, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data['id']
        response = self.api.patch(f'/api/v1/events/{pk}/', {'title':'تحرير الحدث'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.api.delete(f'/api/v1/events/{pk}/').status_code, 204)
        self.assertEqual(HistoricalEvent.objects.get(pk=pk).status, 'ARCHIVED')

    def test_cannot_publish_real_event_without_approved_claim(self):
        self.api.force_authenticate(self.admin)
        response = self.api.patch(f'/api/v1/events/{self.draft.pk}/', {'status':'PUBLISHED'}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_claim_requires_approved_support_before_approval(self):
        claim = HistoricalClaim.objects.create(event=self.event, claim_text='ادعاء بلا دليل', claim_type='FACT')
        self.api.force_authenticate(self.admin)
        response = self.api.patch(f'/api/v1/claims/{claim.pk}/', {'status':'APPROVED'}, format='json')
        self.assertEqual(response.status_code, 400)
        response = self.api.patch(f'/api/v1/claims/{self.claim.pk}/', {'status':'PUBLISHED'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.claim.refresh_from_db()
        self.assertEqual(self.claim.approved_by, self.admin)

    def test_evidence_provenance_must_match_source_and_page(self):
        other = HistoricalSource.objects.create(title='آخر', slug='other-source')
        bad = Evidence(claim=self.claim, source=other, source_chunk=self.chunk, evidence_text=self.claim.claim_text, page_number=7)
        with self.assertRaises(ValidationError): bad.full_clean()
        bad.source = self.source
        bad.page_number = 8
        with self.assertRaises(ValidationError): bad.full_clean()
        bad.page_number = 7
        bad.evidence_text = 'اقتباس مخترع'
        with self.assertRaises(ValidationError): bad.full_clean()

    def test_trace_provenance_and_source_revocation(self):
        data = trace_claim(self.claim)
        self.assertEqual(data['confidence'], 'موثق')
        self.assertEqual(data['evidence'][0]['page'], 7)
        self.source.is_approved = False
        self.source.save()
        data = trace_claim(self.claim)
        self.assertEqual(data['confidence'], 'الأدلة غير كافية')
        self.assertEqual(data['evidence'], [])
        self.assertNotEqual(data['classification'], 'حقيقة موثقة')

    def test_full_admin_publication_journey(self):
        """Exercise upload -> processing -> claim -> evidence -> approval -> publication."""
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.api.force_authenticate(self.admin)
            source = self.api.post('/api/v1/sources/', {'title':'مصدر رحلة اختبار', 'slug':'journey-source', 'file':SimpleUploadedFile('journey.txt', self.text.encode())}, format='multipart')
            self.assertEqual(source.status_code, 201, source.data)
            source_id = source.data['id']
            self.assertEqual(self.api.post(f'/api/v1/sources/{source_id}/process/').status_code, 200)
            self.assertEqual(self.api.patch(f'/api/v1/sources/{source_id}/', {'status':'APPROVED', 'is_approved':True}, format='json').status_code, 200)
            chunk = SourceChunk.objects.get(source_id=source_id)
            event = self.api.post('/api/v1/events/', {'title':'رحلة تحقق', 'slug':'publication-journey', 'era':self.era.pk, 'location':self.place.pk, 'sources':[source_id]}, format='json')
            self.assertEqual(event.status_code, 201, event.data)
            event_id = event.data['id']
            claim = self.api.post('/api/v1/claims/', {'claim_text':self.claim.claim_text, 'claim_type':'FACT', 'event':event_id}, format='json')
            self.assertEqual(claim.status_code, 201, claim.data)
            claim_id = claim.data['id']
            evidence = self.api.post('/api/v1/evidence/', {'claim':claim_id, 'source':source_id, 'source_chunk':chunk.pk, 'evidence_text':self.claim.claim_text, 'section':'نص المصدر'}, format='json')
            self.assertEqual(evidence.status_code, 201, evidence.data)
            self.assertEqual(self.api.patch(f'/api/v1/claims/{claim_id}/', {'status':'APPROVED'}, format='json').status_code, 200)
            self.assertEqual(self.api.patch(f'/api/v1/events/{event_id}/', {'status':'PUBLISHED'}, format='json').status_code, 200)
            self.api.force_authenticate(self.user)
            self.assertEqual(self.api.get(f'/api/v1/events/{event_id}/').status_code, 200)
            trace = self.api.get(f'/api/v1/claims/{claim_id}/evidence/')
            self.assertEqual(trace.data['confidence'], 'موثق')
            answer = self.api.post('/api/v1/ai/ask/', {'event_id':event_id, 'question':'توثيق القرار'}, format='json')
            self.assertEqual(answer.status_code, 200)
            self.assertEqual(answer.data['sources'][0]['id'], source_id)
            self.assertIn(self.claim.claim_text, answer.data['answer'])

    def test_hidden_related_metadata_never_leaks_to_event_html(self):
        self.place.title = 'PRIVATE-PLACE-DO-NOT-LEAK'
        self.place.status = 'DRAFT'
        self.place.save()
        self.era.title = 'PRIVATE-ERA-DO-NOT-LEAK'
        self.era.status = 'DRAFT'
        self.era.save()
        response = self.client.get(self.event.get_absolute_url())
        self.assertNotContains(response, 'PRIVATE-PLACE-DO-NOT-LEAK')
        self.assertNotContains(response, 'PRIVATE-ERA-DO-NOT-LEAK')

    def test_pdf_and_docx_extraction_keep_real_provenance(self):
        from apps.sources.services import extract_document
        from docx import Document
        from reportlab.pdfgen import canvas
        docx_buffer = io.BytesIO()
        document = Document()
        document.add_heading('Test section', level=1)
        document.add_paragraph('Test content with no historical assertions.')
        document.save(docx_buffer)
        pdf_buffer = io.BytesIO()
        pdf = canvas.Canvas(pdf_buffer)
        pdf.drawString(40, 740, 'Test page one.')
        pdf.showPage()
        pdf.drawString(40, 740, 'Test page two.')
        pdf.save()
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            source = HistoricalSource.objects.create(title='DOCX', slug='docx', file=SimpleUploadedFile('test.docx', docx_buffer.getvalue()))
            parts = extract_document(source.file)
            self.assertTrue(any(section == 'Test section' and page is None for page, section, text in parts))
            source = HistoricalSource.objects.create(title='PDF', slug='pdf', file=SimpleUploadedFile('test.pdf', pdf_buffer.getvalue()))
            parts = extract_document(source.file)
            self.assertEqual([page for page, _, _ in parts], [1, 2])
            self.assertIn('Test page two.', parts[1][2])

    def test_ai_rate_limit(self):
        for _ in range(20):
            self.assertEqual(self.api.post('/api/v1/ai/ask/', {'question':'كلمات بدون دليل'}, format='json').status_code, 200)
        self.assertEqual(self.api.post('/api/v1/ai/ask/', {'question':'كلمات بدون دليل'}, format='json').status_code, 429)

    def test_admin_html_editor_saves_and_publishes(self):
        self.client.force_login(self.admin)
        response = self.client.post('/dashboard/persons/new/', {'title':'شخصية تحرير اختبار', 'slug':'editor-person', 'status':'DRAFT'})
        self.assertEqual(response.status_code, 302)
        person = HistoricalPerson.objects.get(slug='editor-person')
        response = self.client.post(f'/dashboard/persons/{person.pk}/', {'title':person.title, 'slug':person.slug, 'status':'APPROVED'})
        self.assertEqual(response.status_code, 302)
        self.client.logout()
        self.assertEqual(self.client.get(person.get_absolute_url()).status_code, 200)

    def test_approved_source_only_retrieval(self):
        chunks, _ = retrieve('توثيق القرار', self.event)
        self.assertEqual([c.pk for c in chunks], [self.chunk.pk])
        self.source.is_approved = False
        self.source.save()
        self.assertEqual(retrieve('توثيق القرار', self.event)[0], [])

    def test_demo_source_never_enters_rag(self):
        self.source.is_demo = True
        self.source.save()
        self.assertEqual(retrieve('توثيق القرار', self.event)[0], [])

    def test_event_scope_does_not_retrieve_unrelated_sources(self):
        self.event.sources.clear()
        self.evidence.delete()
        self.assertEqual(retrieve('توثيق القرار', self.event)[0], [])

    def test_no_evidence_never_returns_confident_fact(self):
        answer = ask('مجرة زحل والمركبة الفضائية', self.user)
        self.assertEqual(answer['answer'], NO_EVIDENCE)
        self.assertEqual(answer['classification'], 'UNCERTAIN')
        self.assertTrue(answer['needs_review'])
        self.assertEqual(answer['sources'], [])
        self.assertTrue(AIQueryLog.objects.filter(user=self.user).exists())

    def test_extractive_mode_returns_real_chunks_and_citations(self):
        answer = ask('توثيق القرار', self.user, self.event)
        self.assertIn(self.text, answer['answer'])
        self.assertEqual(answer['sources'][0]['page'], 7)
        self.assertEqual(answer['mode'], 'extractive')
        self.assertNotEqual(answer['classification'], 'FACT')

    def test_structured_output_rejects_unknown_citations(self):
        for raw in ['{}', 'not json', '{"segments":[{"text":"ادعاء","chunk_ids":[999]}]}', '{"segments":[{"text":"ادعاء","chunk_ids":[]}]}', '{"segments":[{"text":"ادعاء","chunk_ids":[true]}]}']:
            with self.assertRaises((ValueError, TypeError)):
                parse_answer(raw, {self.chunk.pk})

    def test_configured_provider_is_called_and_validated(self):
        with patch('apps.ai.services.get_provider') as provider:
            provider.return_value.name = 'configured'
            provider.return_value.generate.return_value = json.dumps({'segments':[{'text':'شرح مستند إلى المقطع.', 'chunk_ids':[self.chunk.pk]}]})
            answer = ask('توثيق القرار', self.user, self.event)
            provider.return_value.generate.assert_called_once()
        self.assertEqual(answer['mode'], 'configured')
        self.assertTrue(answer['needs_review'])
        self.assertEqual(answer['classification'], 'INTERPRETATION')

    def test_invalid_provider_answer_fails_closed(self):
        with patch('apps.ai.services.get_provider') as provider:
            provider.return_value.generate.return_value = '{"segments":[{"text":"خطأ","chunk_ids":[99999]}]}'
            answer = ask('توثيق القرار', self.user, self.event)
        self.assertEqual(answer['mode'], 'unavailable')
        self.assertEqual(answer['classification'], 'UNCERTAIN')
        self.assertEqual(answer['sources'], [])

    def test_provider_failure_is_graceful(self):
        with patch('apps.ai.services.get_provider', side_effect=TimeoutError('private-secret')):
            answer = ask('توثيق القرار', self.user, self.event)
        self.assertIn('تعذر الوصول', answer['answer'])
        self.assertNotIn('private-secret', json.dumps(answer))

    def test_simulation_requires_explicit_ack_and_is_separate(self):
        path = f'/api/v1/simulations/{self.scenario.pk}/run/'
        self.assertEqual(self.api.post(path, {'acknowledged':False}, format='json').status_code, 400)
        before = HistoricalClaim.objects.count()
        response = self.api.post(path, {'acknowledged':True, 'choice':'بديل تعليمي'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['classification'], 'SIMULATION')
        self.assertIn('ليس حدثًا تاريخيًا', response.data['warning'])
        self.assertEqual(before, HistoricalClaim.objects.count())

    def test_bookmarks_are_private_and_validate_entity(self):
        self.api.force_authenticate(self.user)
        response = self.api.post('/api/v1/bookmarks/', {'entity':self.event.pk}, format='json')
        self.assertEqual(response.status_code, 201)
        pk = response.data['id']
        self.assertEqual(self.api.post('/api/v1/bookmarks/', {'entity':self.draft.pk}, format='json').status_code, 400)
        self.api.force_authenticate(self.other)
        self.assertEqual(self.api.get('/api/v1/bookmarks/').data['count'], 0)
        self.assertEqual(self.api.delete(f'/api/v1/bookmarks/{pk}/').status_code, 404)

    def test_exploration_history_persists(self):
        self.client.force_login(self.user)
        self.client.get(self.event.get_absolute_url())
        self.assertTrue(ExplorationHistory.objects.filter(user=self.user, entity=self.event).exists())

    def test_search_arabic_normalization_preserves_display(self):
        self.assertEqual(normalize('أَثَر إِلَى'), 'اثر الي')
        result = search_entities('ارض الاختبار')
        self.assertEqual(result[0].title, 'أرض الاختبار')
        self.assertNotIn(self.draft.pk, [r.pk for r in search_entities('مسودة')])

    def test_semantic_search_maps_approved_chunks_to_event(self):
        self.chunk.embedding_reference = 'ollama:unit'
        self.chunk.embedding = [1.0, 0.0]
        self.chunk.save()
        with patch.dict(os.environ, {'EMBEDDING_BACKEND':'ollama'}), patch('apps.ai.embeddings.embedding_service') as service:
            service.return_value.reference = 'ollama:unit'
            service.return_value.embed.return_value = [1.0, 0.0]
            ids = [r.pk for r in search_entities('عبارة دلالية بعيدة')]
        self.assertIn(self.event.pk, ids)
        self.assertIn(self.source.pk, ids)
        self.assertNotIn(self.draft.pk, ids)

    def test_graph_uses_db_relations_and_hides_drafts(self):
        EntityRelationship.objects.create(from_entity=self.event, to_entity=self.place, relation_type='OCCURRED_AT', status='PUBLISHED')
        EntityRelationship.objects.create(from_entity=self.event, to_entity=self.draft, relation_type='RELATED_TO', status='PUBLISHED')
        graph = graph_for(self.event)
        self.assertEqual(len(graph['edges']), 1)
        self.assertNotIn(self.draft.pk, [n['id'] for n in graph['nodes']])

    def test_upload_rejects_disguised_files_and_oversize(self):
        for filename, content in [('x.html', b'<html>'), ('x.pdf', b'not a PDF'), ('x.docx', b'not a zip'), ('x.txt', b'abc\x00')]:
            with self.assertRaises(ValidationError): validate_document(SimpleUploadedFile(filename, content))
        big = SimpleUploadedFile('x.txt', b'a')
        big.size = 11 * 1024 * 1024
        with self.assertRaises(ValidationError): validate_document(big)

    def test_txt_ingestion_and_reprocessing_protection(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            source = HistoricalSource.objects.create(title='مصدر اختبار', slug='upload-test', file=SimpleUploadedFile('test.txt', self.text.encode()))
            count = process_source(source)
            self.assertEqual(count, 1)
            chunk = source.chunks.get()
            self.assertIsNone(chunk.page_number)
            self.assertEqual(chunk.text, self.text)
            self.assertEqual(len(chunk.embedding), 256)
            self.assertFalse(source.is_approved)
            Evidence.objects.create(claim=self.claim, source=source, source_chunk=chunk, evidence_text=self.text)
            self.assertEqual(process_source(source),1)
            self.assertEqual(source.chunks.get().pk,chunk.pk)
            self.assertEqual(source.chunks.get().text,self.text)

    def test_source_crud_api_and_processing(self):
        with tempfile.TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            self.api.force_authenticate(self.admin)
            response = self.api.post('/api/v1/sources/', {'title':'رفع اختبار', 'slug':'api-upload', 'file':SimpleUploadedFile('sample.txt', self.text.encode())}, format='multipart')
            self.assertEqual(response.status_code, 201, response.data)
            pk = response.data['id']
            response = self.api.post(f'/api/v1/sources/{pk}/process/', {}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            response = self.api.patch(f'/api/v1/sources/{pk}/', {'status':'APPROVED', 'is_approved':True}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(self.api.delete(f'/api/v1/sources/{pk}/').status_code, 204)
            self.assertFalse(HistoricalSource.objects.get(pk=pk).is_approved)

    def test_csrf_protects_auth_ai_and_anonymous_simulation(self):
        browser = Client(enforce_csrf_checks=True)
        for path, payload in [('/api/v1/auth/login/',{}), ('/api/v1/ai/ask/',{'question':'توثيق القرار'}), (f'/api/v1/simulations/{self.scenario.pk}/run/',{'acknowledged':True})]:
            self.assertEqual(browser.post(path, json.dumps(payload), content_type='application/json').status_code, 403)
        browser.get('/')
        response = browser.post('/api/v1/ai/ask/', json.dumps({'question':'توثيق القرار'}), content_type='application/json', HTTP_X_CSRFTOKEN=browser.cookies['csrftoken'].value)
        self.assertEqual(response.status_code, 200)

    def test_arabic_templates_render_and_escape_untrusted_text(self):
        self.event.description = '<script>alert(1)</script>'
        self.event.save()
        response = self.client.get(self.event.get_absolute_url())
        self.assertContains(response, 'dir="rtl"')
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertContains(response, '&lt;script&gt;')
        self.client.force_login(self.admin)
        for path in ['/dashboard/', '/dashboard/events/', '/dashboard/events/new/', f'/dashboard/events/{self.event.pk}/', '/dashboard/sources/new/', '/dashboard/claims/new/', '/dashboard/evidence/new/', '/dashboard/relationships/new/', '/dashboard/settings/new/']:
            self.assertEqual(self.client.get(path, follow=True).status_code, 200, path)

    def test_seed_idempotent_and_never_creates_approved_sources(self):
        call_command('seed_demo', stdout=io.StringIO())
        count = Entity.objects.count()
        call_command('seed_demo', stdout=io.StringIO())
        self.assertEqual(Entity.objects.count(), count)
        self.assertFalse(HistoricalSource.objects.filter(is_demo=True, is_approved=True).exists())

    @override_settings(DEMO_MODE=False)
    def test_demo_mode_can_be_disabled(self):
        self.event.is_demo = True
        self.event.save()
        self.assertEqual(self.client.get(self.event.get_absolute_url()).status_code, 404)
