import gzip
import io
import json
import os
import socket
import tempfile
import uuid
from unittest.mock import Mock, patch
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, SimpleTestCase, Client, override_settings
from apps.accounts.models import User
from apps.ai.providers import ProviderUnavailable
from apps.dashboard.extraction import generate_suggestions, parse_suggestions, save_drafts
from apps.dashboard.forms import editor_form
from apps.dashboard.models import ExtractionRun
from apps.history.models import HistoricalEra, HistoricalEvent, HistoricalPerson, StoryPassage, Status
from apps.history.story import public_passages, validate_citations
from apps.knowledge.models import HistoricalClaim, Evidence, EntityRelationship
from apps.sources.models import HistoricalSource, SourceChunk
from apps.sources.web_import import public_target, fetch_page, PageText, PinnedHTTPS


class StoryTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.era = HistoricalEra.objects.create(title='حقبة', slug='era', status='PUBLISHED')
        cls.event = HistoricalEvent.objects.create(title='حدث', slug='event', era=cls.era, status='PUBLISHED')
        cls.source = HistoricalSource.objects.create(title='مصدر', slug='source', status='PUBLISHED', is_approved=True)
        cls.chunk = SourceChunk.objects.create(source=cls.source, text='عمل الرجال معا في الحفر.', section='الفصل الأول', embedding_reference='test')
        cls.claim = HistoricalClaim.objects.create(event=cls.event, claim_text='عمل الرجال معا.', status='APPROVED')
        Evidence.objects.create(claim=cls.claim, source=cls.source, source_chunk=cls.chunk, evidence_text=cls.chunk.text, section=cls.chunk.section)
        cls.passage = StoryPassage.objects.create(event=cls.event, chapter='مشهد', text='اجتمع الرجال وبدأوا العمل.', status='PUBLISHED')
        cls.passage.citations.add(cls.claim)

    def test_every_paragraph_opens_its_own_source_and_locator(self):
        self.assertContains(self.client.get(self.event.get_absolute_url()), f'data-story="{self.passage.pk}"')
        result = self.client.get(f'/api/v1/passages/{self.passage.pk}/citations/').json()
        self.assertEqual(result['text'], self.passage.text)
        ref = result['claims'][0]['evidence'][0]
        self.assertEqual((ref['chunk_id'], ref['source_id'], ref['section'], ref['text']), (self.chunk.pk, self.source.pk, self.chunk.section, self.chunk.text))

    def test_any_unavailable_citation_hides_paragraph_and_journey_text(self):
        self.source.is_approved = False; self.source.save()
        self.assertFalse(public_passages().exists())
        self.assertEqual(self.client.get(f'/api/v1/passages/{self.passage.pk}/citations/').status_code, 404)
        page = self.client.get(self.event.get_absolute_url())
        self.assertNotContains(page, self.passage.text)
        station = self.client.get('/journey/').context['stations'][0]
        self.assertEqual(station['story'], [])
        self.assertEqual(station['narrative'], '')

    def test_wrong_event_is_rejected_and_not_public_even_after_direct_db_write(self):
        another = HistoricalEvent.objects.create(title='آخر', slug='another', era=self.era, status='PUBLISHED')
        self.claim.event = another; self.claim.save()
        with self.assertRaises(ValidationError):
            validate_citations(self.passage, [self.claim])
        self.assertFalse(public_passages().exists())

    def test_editor_validates_new_citations_and_api_links_sources_on_publish(self):
        admin = User.objects.create_user('editor', role='ADMIN')
        self.client.force_login(admin)
        self.passage.status = 'DRAFT'; self.passage.save()
        self.event.sources.clear()
        response = self.client.patch(f'/api/v1/passages/{self.passage.pk}/', json.dumps({'status': 'PUBLISHED', 'citations': [self.claim.pk]}), content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(self.event.sources.filter(pk=self.source.pk).exists())
        response = self.client.patch(f'/api/v1/passages/{self.passage.pk}/', json.dumps({'citations': []}), content_type='application/json')
        self.assertEqual(response.status_code, 400)
        form = editor_form(StoryPassage)({'event':self.event.pk, 'chapter':'مشهد', 'text':'نص', 'position':1, 'status':'PUBLISHED', 'citations':[]}, instance=self.passage)
        self.assertFalse(form.is_valid())

    def test_draft_story_never_leaks_and_html_is_escaped(self):
        self.passage.text = '<script>alert(1)</script>'; self.passage.save()
        page = self.client.get(self.event.get_absolute_url())
        self.assertNotContains(page, self.passage.text)
        self.assertContains(page, '&lt;script&gt;')
        self.passage.status = 'DRAFT'; self.passage.save()
        self.assertEqual(self.client.get(f'/api/v1/passages/{self.passage.pk}/').status_code, 404)


class StoryPackTests(TestCase):
    def test_pack_preview_apply_references_depth_and_preservation(self):
        call_command('import_dorar_pilot', apply=True, stdout=io.StringIO())
        call_command('enrich_medina', apply=True, stdout=io.StringIO())
        before = (HistoricalEvent.objects.count(), HistoricalSource.objects.count())
        call_command('enrich_stories', stdout=io.StringIO())
        self.assertFalse(StoryPassage.objects.exists())
        self.assertEqual(before, (HistoricalEvent.objects.count(), HistoricalSource.objects.count()))
        call_command('enrich_stories', apply=True, stdout=io.StringIO())
        self.assertEqual(public_passages().count(), 24)
        self.assertEqual(HistoricalEvent.objects.count(), before[0])
        for slug in ['badr', 'uhud', 'khandaq']:
            passages = list(public_passages().filter(event__slug=slug))
            self.assertGreaterEqual(sum(len(p.text.split()) for p in passages), 400)
            self.assertGreaterEqual(len({p.chapter for p in passages}), 4)
            for passage in passages:
                validate_citations(passage, passage.citations.all())
                for claim in passage.citations.all():
                    for proof in claim.evidence.all():
                        proof.full_clean()
        first = StoryPassage.objects.first(); first.text = 'تحرير لاحق'; first.save()
        call_command('enrich_stories', apply=True, stdout=io.StringIO())
        first.refresh_from_db(); self.assertEqual(first.text, 'تحرير لاحق')


class StudioTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name)
        self.settings_override.enable(); self.addCleanup(self.settings_override.disable)
        self.admin = User.objects.create_user('librarian', role='ADMIN')
        era = HistoricalEra.objects.create(title='حقبة', slug='era', status='PUBLISHED')
        self.event = HistoricalEvent.objects.create(title='حدث', slug='event', era=era, status='PUBLISHED')
        self.person = HistoricalPerson.objects.create(title='شخصية', slug='person', status='PUBLISHED')
        self.source = HistoricalSource.objects.create(title='مصدر خاص', slug='private')
        self.chunk = SourceChunk.objects.create(source=self.source, text='شارك جابر في العمل. ورد تفصيل في هذا النص.', page_number=7, section='باب العمل', embedding_reference='test')
        self.run = ExtractionRun.objects.create(event=self.event, source=self.source, created_by=self.admin)
        self.row = {'claim':'شارك جابر في العمل.', 'classification':'FACT', 'chapter':'عمل الجماعة', 'story':'شارك جابر في العمل مع الجماعة.',
            'person_id':self.person.pk, 'place_id':None, 'role':'PARTICIPATED_IN', 'citations':[{'chunk_id':self.chunk.pk,'quote':'شارك جابر في العمل.'}]}

    def generate(self):
        provider = Mock(); provider.generate.return_value = json.dumps({'items':[self.row]})
        with patch('apps.dashboard.extraction.get_provider', return_value=provider):
            generate_suggestions(self.run, [self.chunk.pk])
        return provider

    def test_admin_only_and_csrf(self):
        self.assertEqual(self.client.get('/dashboard/studio/').status_code, 302)
        member = User.objects.create_user('member')
        self.client.force_login(member)
        for path in ['/dashboard/studio/', f'/dashboard/studio/{self.run.pk}/']:
            self.assertEqual(self.client.get(path).status_code, 403)
        secure = Client(enforce_csrf_checks=True); secure.force_login(self.admin)
        self.assertEqual(secure.post(f'/dashboard/studio/{self.run.pk}/extract/', {'chunks':[self.chunk.pk]}).status_code, 403)

    def test_txt_import_is_private_has_locators_and_token_is_idempotent(self):
        self.client.force_login(self.admin)
        token = str(uuid.uuid4())
        def upload():
            return self.client.post('/dashboard/studio/', {'token':token,'event':self.event.pk,'title':'ملف للقراءة','input_kind':'file',
                'document':SimpleUploadedFile('history.txt', 'نص عربي موثق للمراجعة.'.encode('utf-8'))})
        response = upload(); self.assertEqual(response.status_code, 302)
        run = ExtractionRun.objects.get(token=token)
        self.assertEqual(run.source.status, 'DRAFT'); self.assertFalse(run.source.is_approved)
        self.assertEqual(run.source.chunks.get().section, 'نص المصدر')
        self.assertEqual(self.client.get(run.source.get_absolute_url()).status_code, 404)
        self.assertEqual(upload().status_code, 302)
        self.assertEqual(ExtractionRun.objects.filter(token=token).count(), 1)

    def test_docx_upload_preserves_heading(self):
        from docx import Document
        document = Document(); document.add_heading('فصل المشورة', 1); document.add_paragraph('اجتمعت الجماعة للمشورة.')
        blob = io.BytesIO(); document.save(blob)
        self.client.force_login(self.admin)
        token = uuid.uuid4()
        response = self.client.post('/dashboard/studio/', {'token':str(token), 'event':self.event.pk, 'title':'وثيقة', 'input_kind':'file',
            'document':SimpleUploadedFile('history.docx', blob.getvalue())})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ExtractionRun.objects.get(token=token).source.chunks.get().section, 'فصل المشورة')

    def test_pdf_upload_preserves_actual_page_number(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 260 Td (Historical source for review.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        blob = io.BytesIO(); writer.write(blob)
        self.client.force_login(self.admin); token = uuid.uuid4()
        response = self.client.post('/dashboard/studio/', {'token':str(token), 'event':self.event.pk, 'title':'PDF', 'input_kind':'file',
            'document':SimpleUploadedFile('history.pdf', blob.getvalue())})
        self.assertEqual(response.status_code, 302)
        chunk = ExtractionRun.objects.get(token=token).source.chunks.get()
        self.assertEqual(chunk.page_number, 1); self.assertIn('Historical source for review.', chunk.text)

    def test_url_import_stores_final_url_but_not_published(self):
        self.client.force_login(self.admin)
        token = uuid.uuid4()
        with patch('apps.dashboard.import_views.fetch_page', return_value=('https://example.org/final', 'مادة المصدر للقراءة والتحليل.')):
            response = self.client.post('/dashboard/studio/', {'token':str(token), 'event':self.event.pk, 'title':'موقع', 'input_kind':'url', 'url':'https://example.org/start'})
        self.assertEqual(response.status_code, 302)
        source = ExtractionRun.objects.get(token=token).source
        self.assertEqual(source.url, 'https://example.org/final'); self.assertFalse(source.is_approved)
        self.assertEqual(source.chunks.get().text, 'مادة المصدر للقراءة والتحليل.')

    def test_upload_error_rolls_back_and_cleans_only_new_file(self):
        self.client.force_login(self.admin)
        before = HistoricalSource.objects.count()
        response = self.client.post('/dashboard/studio/', {'token':str(uuid.uuid4()), 'event':self.event.pk,'title':'فارغ','input_kind':'file',
            'document':SimpleUploadedFile('empty.txt', b'   ')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(HistoricalSource.objects.count(), before)
        from pathlib import Path
        self.assertEqual([p for p in Path(self.media.name).rglob('*') if p.is_file()], [])

    def test_generation_only_sends_selected_source_chunks_then_review_saves_drafts(self):
        other = SourceChunk.objects.create(source=self.source, ordinal=1, text='نص لم يتم اختياره', embedding_reference='test')
        provider = self.generate()
        self.assertEqual(self.run.status, 'REVIEW')
        sent = provider.generate.call_args.args[1]
        self.assertEqual([c['id'] for c in sent['chunks']], [self.chunk.pk])
        self.assertEqual(sent['max_items'],1)
        self.assertIn('schema',provider.generate.call_args.kwargs)
        self.assertEqual(HistoricalClaim.objects.count(), 0)
        saved = save_drafts(self.run, {0:{'story':'صياغة حررها المسؤول.'}}, self.admin)
        claim = HistoricalClaim.objects.get(pk=saved.created_claim_ids[0]); passage = StoryPassage.objects.get(pk=saved.created_passage_ids[0])
        self.assertEqual(claim.status, 'DRAFT'); self.assertEqual(passage.status, 'DRAFT')
        self.assertEqual(passage.text, 'صياغة حررها المسؤول.'); self.assertTrue(passage.ai_suggested)
        self.assertEqual(claim.evidence.get().page_number, 7)
        self.assertEqual(passage.citations.get(), claim)
        edge = EntityRelationship.objects.get(); self.assertEqual(edge.supporting_claim_id, claim.pk); self.assertEqual(edge.status, 'DRAFT')
        self.assertFalse(public_passages().exists())
        with self.assertRaises(ValidationError):
            save_drafts(saved, {0:{}}, self.admin)
        self.assertEqual(HistoricalClaim.objects.count(), 1)

    def test_review_post_keeps_editor_changes_and_unselected_items_are_ignored(self):
        self.generate(); self.client.force_login(self.admin)
        response = self.client.post(f'/dashboard/studio/{self.run.pk}/', {'0-selected':'on','0-claim':'معلومة راجعها المحرر.',
            '0-story':'سرد راجعه المحرر.','0-chapter':'فصل جديد','0-classification':'INTERPRETATION','0-person':self.person.pk,'0-role':'NARRATED'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(StoryPassage.objects.get().text, 'سرد راجعه المحرر.')
        self.assertEqual(EntityRelationship.objects.get().relation_type, 'NARRATED')

    def test_hallucinated_quotes_ids_and_entities_are_rejected(self):
        for row in [{**self.row,'citations':[{'chunk_id':999999,'quote':'نص'}]}, {**self.row,'citations':[{'chunk_id':self.chunk.pk,'quote':'اقتباس مخترع'}]}, {**self.row,'person_id':999999}]:
            with self.subTest(row=row), self.assertRaises(ValidationError):
                parse_suggestions(json.dumps({'items':[row]}), {self.chunk.pk:self.chunk}, {self.person.pk}, set())

    def test_cross_source_chunk_selection_is_rejected_before_provider(self):
        other_source = HistoricalSource.objects.create(title='آخر', slug='other')
        chunk = SourceChunk.objects.create(source=other_source, text='نص', embedding_reference='test')
        with patch('apps.dashboard.extraction.get_provider') as provider, self.assertRaises(ValidationError):
            generate_suggestions(self.run, [chunk.pk])
        provider.assert_not_called()

    def test_provider_denial_is_honest_retryable_and_keeps_source(self):
        provider = Mock(); provider.generate.side_effect = ProviderUnavailable('project_access_denied')
        with patch('apps.dashboard.extraction.get_provider', return_value=provider):
            generate_suggestions(self.run, [self.chunk.pk])
        self.assertEqual(self.run.status, 'FAILED'); self.assertIn('403', self.run.error_message)
        self.assertFalse(StoryPassage.objects.exists()); self.assertTrue(SourceChunk.objects.filter(pk=self.chunk.pk).exists())
        self.generate(); self.assertEqual(self.run.status, 'REVIEW')

    def test_modified_source_quote_and_mixed_invalid_review_roll_back_whole_save(self):
        self.generate(); self.run.suggestions.append({**self.row,'citations':[{'chunk_id':self.chunk.pk,'quote':'غير موجود'}]}); self.run.save()
        with self.assertRaises(ValidationError):
            save_drafts(self.run, {0:{},1:{}}, self.admin)
        self.assertFalse(StoryPassage.objects.exists()); self.assertFalse(HistoricalClaim.objects.exists())
        self.run.refresh_from_db(); self.assertEqual(self.run.status, 'REVIEW')


class WebImportTests(SimpleTestCase):
    def addresses(self, *addresses):
        return [(socket.AF_INET6 if ':' in ip else socket.AF_INET, socket.SOCK_STREAM, 6, '', (ip,443)) for ip in addresses]

    def test_private_ipv4_ipv6_mixed_dns_and_credentials_are_rejected(self):
        for ips in [('127.0.0.1',),('169.254.169.254',),('::1',),('fc00::1',),('93.184.216.34','10.0.0.1')]:
            with self.subTest(ips=ips), patch('apps.sources.web_import.socket.getaddrinfo', return_value=self.addresses(*ips)), self.assertRaises(ValidationError):
                public_target('https://example.org/page')
        for url in ['file:///etc/passwd','http://user:pass@example.org/','https://example.org:8443/','http://example.org\\@127.0.0.1/']:
            with self.subTest(url=url), self.assertRaises(ValidationError):
                public_target(url)

    def test_connection_pins_validated_ip_and_preserves_tls_hostname(self):
        connection = PinnedHTTPS('example.org', 443, '93.184.216.34')
        context = Mock(); connection._context = context
        with patch('apps.sources.web_import.socket.create_connection') as connect:
            connection.connect()
            connect.assert_called_once_with(('93.184.216.34',443),10)
            context.wrap_socket.assert_called_once_with(connect.return_value, server_hostname='example.org')

    def test_redirect_target_is_revalidated(self):
        response = Mock(status=302); response.getheader.return_value = 'http://127.0.0.1/'
        connection = Mock(); connection.getresponse.return_value = response
        with patch('apps.sources.web_import.socket.getaddrinfo', side_effect=[self.addresses('93.184.216.34'),self.addresses('127.0.0.1')]), patch('apps.sources.web_import.PinnedHTTPS', return_value=connection), self.assertRaises(ValidationError):
            fetch_page('https://example.org/')
        connection.close.assert_called_once()

    def test_html_excludes_scripts_and_navigation(self):
        parser = PageText(); parser.feed('<nav>navigation</nav><main><h1>عنوان</h1><p>' + 'نص تاريخي. '*20 + '</p><script>steal()</script></main>')
        self.assertNotIn('steal', parser.text()); self.assertNotIn('navigation',parser.text()); self.assertIn('عنوان', parser.text())

    def test_download_and_decompression_are_bounded(self):
        for headers, body in [({'Content-Type':'application/pdf'},b'pdf'),({'Content-Type':'text/html'},b'x'*300),
                              ({'Content-Type':'text/html','Content-Encoding':'gzip'},gzip.compress(b'x'*500))]:
            response = Mock(status=200); response.getheader.side_effect=lambda key, default=None: headers.get(key,default)
            stream = io.BytesIO(body); response.read.side_effect = stream.read
            connection = Mock(); connection.getresponse.return_value = response
            with self.subTest(headers=headers), patch('apps.sources.web_import.MAX_BYTES', 200), patch('apps.sources.web_import.public_target', return_value=(__import__('urllib.parse',fromlist=['urlsplit']).urlsplit('https://example.org/'),'example.org',443,'93.184.216.34')), patch('apps.sources.web_import.PinnedHTTPS', return_value=connection), self.assertRaises(ValidationError):
                fetch_page('https://example.org/')
